"""
Export Service - Xuất PDF và Excel cho thông báo
"""
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models.notification import Notification


class ExportService:
    """Service để xuất báo cáo PDF và Excel"""

    @staticmethod
    def export_notifications_to_excel(
        notifications: list[Notification],
        parent_name: str,
        export_options: dict
    ) -> BytesIO:
        """
        Xuất danh sách thông báo ra Excel

        Args:
            notifications: Danh sách thông báo
            parent_name: Tên phụ huynh
            export_options: Tùy chọn xuất (title, content, marks, etc.)

        Returns:
            BytesIO: File Excel
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Thông Báo"

        # Styling
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        cell_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Title
        ws.merge_cells('A1:F1')
        ws['A1'] = 'DANH SÁCH THÔNG BÁO'
        ws['A1'].font = Font(bold=True, size=16)
        ws['A1'].alignment = Alignment(horizontal='center')

        # Info
        ws['A2'] = f'Phụ huynh: {parent_name}'
        ws['A3'] = f'Ngày xuất: {datetime.now().strftime("%d/%m/%Y %H:%M")}'
        ws['A4'] = f'Tổng số thông báo: {len(notifications)}'

        # Headers
        row = 6
        col = 1
        headers = []

        if export_options.get('title', True):
            headers.append('Tiêu Đề')
        if export_options.get('content', True):
            headers.append('Nội Dung')
        if export_options.get('marks', True):
            headers.append('Loại')
        if export_options.get('sender', True):
            headers.append('Trạng Thái')
        if export_options.get('time', True):
            headers.append('Thời Gian')

        # Write headers
        for idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=idx)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border

        # Write data
        for notif in notifications:
            row += 1
            col = 1

            if export_options.get('title', True):
                cell = ws.cell(row=row, column=col)
                cell.value = notif.title
                cell.alignment = cell_alignment
                cell.border = border
                col += 1

            if export_options.get('content', True):
                cell = ws.cell(row=row, column=col)
                cell.value = notif.message
                cell.alignment = cell_alignment
                cell.border = border
                col += 1

            if export_options.get('marks', True):
                cell = ws.cell(row=row, column=col)
                type_map = {
                    'grade': 'Điểm số',
                    'success': 'Hoàn thành',
                    'warning': 'Cảnh báo',
                    'alert': 'Khẩn cấp',
                    'info': 'Thông tin'
                }
                cell.value = type_map.get(notif.type, notif.type)
                cell.alignment = cell_alignment
                cell.border = border
                col += 1

            if export_options.get('sender', True):
                cell = ws.cell(row=row, column=col)
                cell.value = 'Đã đọc' if notif.is_read else 'Chưa đọc'
                cell.alignment = cell_alignment
                cell.border = border

                # Color code
                if notif.is_read:
                    cell.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
                else:
                    cell.fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

                col += 1

            if export_options.get('time', True):
                cell = ws.cell(row=row, column=col)
                cell.value = notif.created_at.strftime("%d/%m/%Y %H:%M")
                cell.alignment = cell_alignment
                cell.border = border
                col += 1

        # Adjust column widths
        column_widths = {
            1: 30,  # Tiêu đề
            2: 50,  # Nội dung
            3: 15,  # Loại
            4: 15,  # Trạng thái
            5: 20   # Thời gian
        }

        for col_num, width in column_widths.items():
            if col_num <= len(headers):
                ws.column_dimensions[chr(64 + col_num)].width = width

        # Set row heights
        for row_num in range(7, row + 1):
            ws.row_dimensions[row_num].height = 30

        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return output

    @staticmethod
    def export_notifications_to_pdf(
        notifications: list[Notification],
        parent_name: str,
        export_options: dict
    ) -> BytesIO:
        """
        Xuất danh sách thông báo ra PDF

        Args:
            notifications: Danh sách thông báo
            parent_name: Tên phụ huynh
            export_options: Tùy chọn xuất

        Returns:
            BytesIO: File PDF
        """
        buffer = BytesIO()

        # Create PDF
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=2*cm,
            leftMargin=2*cm,
            topMargin=2*cm,
            bottomMargin=2*cm
        )

        # Container for elements
        elements = []

        # Styles
        styles = getSampleStyleSheet()

        # Title style
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1f2937'),
            spaceAfter=20,
            alignment=1  # Center
        )

        # Heading style
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=14,
            textColor=colors.HexColor('#4472C4'),
            spaceAfter=10
        )

        # Normal style
        normal_style = ParagraphStyle(
            'CustomNormal',
            parent=styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#374151')
        )

        # Title
        title = Paragraph("DANH SÁCH THÔNG BÁO", title_style)
        elements.append(title)
        elements.append(Spacer(1, 0.5*cm))

        # Info
        info_data = [
            ['Phụ huynh:', parent_name],
            ['Ngày xuất:', datetime.now().strftime("%d/%m/%Y %H:%M")],
            ['Tổng số thông báo:', str(len(notifications))]
        ]

        info_table = Table(info_data, colWidths=[4*cm, 12*cm])
        info_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#6b7280')),
            ('TEXTCOLOR', (1, 0), (1, -1), colors.HexColor('#1f2937')),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ]))

        elements.append(info_table)
        elements.append(Spacer(1, 1*cm))

        # Notifications
        for idx, notif in enumerate(notifications, start=1):
            # Notification header
            type_map = {
                'grade': 'Điểm số',
                'success': 'Hoàn thành',
                'warning': 'Cảnh báo',
                'alert': 'Khẩn cấp',
                'info': 'Thông tin'
            }

            type_text = type_map.get(notif.type, notif.type)
            status_text = 'Đã đọc' if notif.is_read else 'Chưa đọc'
            time_text = notif.created_at.strftime("%d/%m/%Y %H:%M")

            # Header row
            header_text = f"<b>{idx}. {notif.title}</b>"
            header_para = Paragraph(header_text, heading_style)
            elements.append(header_para)

            # Meta info
            meta_data = []

            if export_options.get('marks', True):
                meta_data.append(['Loại:', type_text])

            if export_options.get('sender', True):
                meta_data.append(['Trạng thái:', status_text])

            if export_options.get('time', True):
                meta_data.append(['Thời gian:', time_text])

            if meta_data:
                meta_table = Table(meta_data, colWidths=[3*cm, 13*cm])
                meta_table.setStyle(TableStyle([
                    ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 0), (-1, -1), 9),
                    ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#6b7280')),
                    ('TEXTCOLOR', (1, 0), (1, -1), colors.HexColor('#374151')),
                    ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ]))
                elements.append(meta_table)
                elements.append(Spacer(1, 0.3*cm))

            # Content
            if export_options.get('content', True):
                content_para = Paragraph(f"<i>{notif.message}</i>", normal_style)

                content_table = Table([[content_para]], colWidths=[16*cm])
                content_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f9fafb')),
                    ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#e5e7eb')),
                    ('LEFTPADDING', (0, 0), (-1, -1), 10),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 10),
                    ('TOPPADDING', (0, 0), (-1, -1), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
                ]))

                elements.append(content_table)

            elements.append(Spacer(1, 0.7*cm))

            # Page break after every 3 notifications (except last)
            if idx % 3 == 0 and idx < len(notifications):
                elements.append(PageBreak())

        # Build PDF
        doc.build(elements)
        buffer.seek(0)

        return buffer

    @staticmethod
    def get_filename(parent_name: str, file_type: str) -> str:
        """
        Tạo tên file cho export

        Args:
            parent_name: Tên phụ huynh
            file_type: 'pdf' hoặc 'excel'

        Returns:
            str: Tên file
        """
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        safe_name = parent_name.replace(' ', '_')

        if file_type == 'pdf':
            return f"ThongBao_{safe_name}_{timestamp}.pdf"
        if file_type == 'excel':
            return f"ThongBao_{safe_name}_{timestamp}.xlsx"
        return f"ThongBao_{safe_name}_{timestamp}.{file_type}"


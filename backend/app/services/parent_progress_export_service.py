from io import BytesIO
from datetime import datetime
from typing import List, Dict, Any
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


class ParentProgressExportService:
    """Service for exporting parent progress reports to PDF and Excel"""
    
    @staticmethod
    def export_to_pdf(
        student_info: Dict[str, Any],
        progress_data: Dict[str, Any],
        export_options: Dict[str, bool]
    ) -> BytesIO:
        """Export parent progress report to PDF"""
        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            rightMargin=72,
            leftMargin=72,
            topMargin=72,
            bottomMargin=18,
        )
        
        # Container for the 'Flowable' objects
        elements = []
        
        # Define styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#6366f1'),
            spaceAfter=30,
            alignment=TA_CENTER,
            fontName='Helvetica-Bold'
        )
        
        heading_style = ParagraphStyle(
            'CustomHeading',
            parent=styles['Heading2'],
            fontSize=16,
            textColor=colors.HexColor('#4f46e5'),
            spaceAfter=12,
            spaceBefore=12,
            fontName='Helvetica-Bold'
        )
        
        normal_style = styles['Normal']
        normal_style.fontSize = 10
        
        # Title
        title = Paragraph("BÁO CÁO TIẾN ĐỘ HỌC TẬP", title_style)
        elements.append(title)
        elements.append(Spacer(1, 12))
        
        # Student Info Section
        if export_options.get('studentInfo', True):
            elements.append(Paragraph("THÔNG TIN HỌC SINH", heading_style))
            
            student_data = [
                ['Họ và tên:', student_info.get('name', 'N/A')],
                ['Email:', student_info.get('email', 'N/A')],
                ['Lớp:', student_info.get('grade', 'N/A')],
                ['Ngày xuất báo cáo:', datetime.now().strftime('%d/%m/%Y %H:%M')],
            ]
            
            student_table = Table(student_data, colWidths=[2*inch, 4*inch])
            student_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
                ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#1f2937')),
                ('ALIGN', (0, 0), (0, -1), 'LEFT'),
                ('ALIGN', (1, 0), (1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb'))
            ]))
            
            elements.append(student_table)
            elements.append(Spacer(1, 20))
        
        # Overall Statistics
        elements.append(Paragraph("THỐNG KÊ TỔNG QUAN", heading_style))
        
        overall_data = [
            ['Điểm trung bình:', f"{progress_data.get('overall_average', 0):.1f}/10"],
            ['Tổng số bài nộp:', str(progress_data.get('total_submissions', 0))],
            ['Tổng số buổi học:', str(progress_data.get('attendance', {}).get('total', 0))],
            ['Số buổi có mặt:', str(progress_data.get('attendance', {}).get('present', 0))],
        ]
        
        overall_table = Table(overall_data, colWidths=[2*inch, 4*inch])
        overall_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#f3f4f6')),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#1f2937')),
            ('ALIGN', (0, 0), (0, -1), 'LEFT'),
            ('ALIGN', (1, 0), (1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb'))
        ]))
        
        elements.append(overall_table)
        elements.append(Spacer(1, 20))
        
        # Progress by Subject
        if export_options.get('progressChart', True):
            elements.append(Paragraph("TIẾN ĐỘ THEO KỸ NĂNG", heading_style))
            
            subject_progress = progress_data.get('subject_progress', [])
            if subject_progress:
                subject_data = [['Kỹ năng', 'Hoàn thành', 'Điểm TB', 'Số bài']]
                for subject in subject_progress:
                    subject_data.append([
                        subject.get('subject', ''),
                        f"{subject.get('progress', 0)}%",
                        f"{subject.get('average_score', 0) or 'N/A'}",
                        f"{subject.get('completed_exercises', 0)}/{subject.get('total_exercises', 0)}"
                    ])
                
                subject_table = Table(subject_data, colWidths=[1.5*inch, 1.5*inch, 1.5*inch, 1.5*inch])
                subject_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#6366f1')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 11),
                    ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 1), (-1, -1), 10),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                    ('TOPPADDING', (0, 0), (-1, -1), 8),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9fafb')])
                ]))
                
                elements.append(subject_table)
                elements.append(Spacer(1, 20))
        
        # Attendance
        if export_options.get('attendance', True):
            elements.append(Paragraph("CHUYÊN CẦN", heading_style))
            
            attendance = progress_data.get('attendance', {})
            attendance_data = [
                ['Buổi có mặt', 'Buổi vắng', 'Buổi đi muộn', 'Tổng số buổi'],
                [
                    str(attendance.get('present', 0)),
                    str(attendance.get('absent', 0)),
                    str(attendance.get('late', 0)),
                    str(attendance.get('total', 0))
                ]
            ]
            
            attendance_table = Table(attendance_data, colWidths=[1.5*inch, 1.5*inch, 1.5*inch, 1.5*inch])
            attendance_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#10b981')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 11),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb'))
            ]))
            
            elements.append(attendance_table)
            elements.append(Spacer(1, 20))
        
        # Detailed Grades
        if export_options.get('grades', True):
            elements.append(PageBreak())
            elements.append(Paragraph("CHI TIẾT ĐIỂM SỐ", heading_style))
            
            grades = progress_data.get('detailed_grades', [])[:20]  # Limit to 20 for PDF
            if grades:
                grades_data = [['Bài tập', 'Kỹ năng', 'Điểm', 'Ngày nộp']]
                for grade in grades:
                    grades_data.append([
                        Paragraph(grade.get('exercise_title', '')[:40], normal_style),
                        grade.get('skill_type', 'N/A') or 'Mixed',
                        f"{grade.get('score', 'N/A') or 'N/A'}",
                        grade.get('submitted_at', 'N/A')[:10] if grade.get('submitted_at') else 'N/A'
                    ])
                
                grades_table = Table(grades_data, colWidths=[3*inch, 1*inch, 0.8*inch, 1.2*inch])
                grades_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f59e0b')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (0, -1), 'LEFT'),
                    ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 10),
                    ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                    ('FONTSIZE', (0, 1), (-1, -1), 9),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 6),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb')),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#fffbeb')])
                ]))
                
                elements.append(grades_table)
                elements.append(Spacer(1, 20))
        
        # Teacher Comments
        if export_options.get('teacherComments', True):
            elements.append(Paragraph("NHẬN XÉT CỦA GIÁO VIÊN", heading_style))
            
            grades_with_feedback = [g for g in progress_data.get('detailed_grades', []) if g.get('feedback')][:10]
            if grades_with_feedback:
                for grade in grades_with_feedback:
                    comment_data = [
                        [Paragraph(f"<b>{grade.get('exercise_title', '')}</b>", normal_style)],
                        [Paragraph(f"<i>{grade.get('feedback', 'Không có nhận xét')}</i>", normal_style)]
                    ]
                    
                    comment_table = Table(comment_data, colWidths=[6*inch])
                    comment_table.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f3f4f6')),
                        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
                        ('TEXTCOLOR', (0, 0), (-1, -1), colors.HexColor('#1f2937')),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('FONTSIZE', (0, 0), (-1, -1), 9),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                        ('TOPPADDING', (0, 0), (-1, -1), 6),
                        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb'))
                    ]))
                    
                    elements.append(comment_table)
                    elements.append(Spacer(1, 10))
            else:
                elements.append(Paragraph("Chưa có nhận xét từ giáo viên.", normal_style))
                elements.append(Spacer(1, 10))
        
        # Overall Evaluation
        if export_options.get('overallEvaluation', True):
            elements.append(Spacer(1, 20))
            elements.append(Paragraph("ĐÁNH GIÁ TỔNG QUAN", heading_style))
            
            overall_avg = progress_data.get('overall_average', 0)
            evaluation_text = ""
            if overall_avg >= 8.5:
                evaluation_text = "Học sinh có thành tích xuất sắc, tiếp tục phát huy."
            elif overall_avg >= 7.0:
                evaluation_text = "Học sinh có thành tích khá tốt, cần tiếp tục cố gắng."
            elif overall_avg >= 5.0:
                evaluation_text = "Học sinh đạt thành tích trung bình, cần nỗ lực hơn."
            else:
                evaluation_text = "Học sinh cần cải thiện kết quả học tập, cần sự hỗ trợ thêm từ gia đình và thầy cô."
            
            elements.append(Paragraph(evaluation_text, normal_style))
        
        # Build PDF
        doc.build(elements)
        buffer.seek(0)
        return buffer
    
    @staticmethod
    def export_to_excel(
        student_info: Dict[str, Any],
        progress_data: Dict[str, Any],
        export_options: Dict[str, bool]
    ) -> BytesIO:
        """Export parent progress report to Excel"""
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet
        
        # Define styles
        header_font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='6366F1', end_color='6366F1', fill_type='solid')
        
        title_font = Font(name='Arial', size=16, bold=True, color='4F46E5')
        subtitle_font = Font(name='Arial', size=11, bold=True)
        normal_font = Font(name='Arial', size=10)
        
        center_alignment = Alignment(horizontal='center', vertical='center')
        left_alignment = Alignment(horizontal='left', vertical='center')
        
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Sheet 1: Overview
        ws_overview = wb.create_sheet("Tổng quan")
        
        # Title
        ws_overview['A1'] = 'BÁO CÁO TIẾN ĐỘ HỌC TẬP'
        ws_overview['A1'].font = title_font
        ws_overview['A1'].alignment = center_alignment
        ws_overview.merge_cells('A1:D1')
        
        row = 3
        
        # Student Info
        if export_options.get('studentInfo', True):
            ws_overview[f'A{row}'] = 'THÔNG TIN HỌC SINH'
            ws_overview[f'A{row}'].font = subtitle_font
            row += 1
            
            ws_overview[f'A{row}'] = 'Họ và tên:'
            ws_overview[f'B{row}'] = student_info.get('name', 'N/A')
            row += 1
            
            ws_overview[f'A{row}'] = 'Email:'
            ws_overview[f'B{row}'] = student_info.get('email', 'N/A')
            row += 1
            
            ws_overview[f'A{row}'] = 'Lớp:'
            ws_overview[f'B{row}'] = student_info.get('grade', 'N/A')
            row += 1
            
            ws_overview[f'A{row}'] = 'Ngày xuất báo cáo:'
            ws_overview[f'B{row}'] = datetime.now().strftime('%d/%m/%Y %H:%M')
            row += 2
        
        # Overall Statistics
        ws_overview[f'A{row}'] = 'THỐNG KÊ TỔNG QUAN'
        ws_overview[f'A{row}'].font = subtitle_font
        row += 1
        
        ws_overview[f'A{row}'] = 'Điểm trung bình:'
        ws_overview[f'B{row}'] = f"{progress_data.get('overall_average', 0):.1f}/10"
        row += 1
        
        ws_overview[f'A{row}'] = 'Tổng số bài nộp:'
        ws_overview[f'B{row}'] = progress_data.get('total_submissions', 0)
        row += 1
        
        ws_overview[f'A{row}'] = 'Tổng số buổi học:'
        ws_overview[f'B{row}'] = progress_data.get('attendance', {}).get('total', 0)
        row += 1
        
        ws_overview[f'A{row}'] = 'Số buổi có mặt:'
        ws_overview[f'B{row}'] = progress_data.get('attendance', {}).get('present', 0)
        row += 2
        
        # Progress by Skill
        if export_options.get('progressChart', True):
            ws_overview[f'A{row}'] = 'TIẾN ĐỘ THEO KỸ NĂNG'
            ws_overview[f'A{row}'].font = subtitle_font
            row += 1
            
            # Headers
            headers = ['Kỹ năng', 'Hoàn thành', 'Điểm TB', 'Số bài hoàn thành', 'Tổng số bài']
            for col, header in enumerate(headers, start=1):
                cell = ws_overview.cell(row=row, column=col)
                cell.value = header
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_alignment
                cell.border = thin_border
            row += 1
            
            # Data
            for subject in progress_data.get('subject_progress', []):
                ws_overview[f'A{row}'] = subject.get('subject', '')
                ws_overview[f'B{row}'] = f"{subject.get('progress', 0)}%"
                ws_overview[f'C{row}'] = subject.get('average_score') or 'N/A'
                ws_overview[f'D{row}'] = subject.get('completed_exercises', 0)
                ws_overview[f'E{row}'] = subject.get('total_exercises', 0)
                
                for col in range(1, 6):
                    cell = ws_overview.cell(row=row, column=col)
                    cell.font = normal_font
                    cell.alignment = center_alignment
                    cell.border = thin_border
                
                row += 1
            
            row += 2
        
        # Attendance
        if export_options.get('attendance', True):
            ws_overview[f'A{row}'] = 'CHUYÊN CẦN'
            ws_overview[f'A{row}'].font = subtitle_font
            row += 1
            
            attendance = progress_data.get('attendance', {})
            attendance_headers = ['Buổi có mặt', 'Buổi vắng', 'Buổi đi muộn', 'Tổng số buổi']
            for col, header in enumerate(attendance_headers, start=1):
                cell = ws_overview.cell(row=row, column=col)
                cell.value = header
                cell.font = header_font
                cell.fill = PatternFill(start_color='10B981', end_color='10B981', fill_type='solid')
                cell.alignment = center_alignment
                cell.border = thin_border
            row += 1
            
            ws_overview[f'A{row}'] = attendance.get('present', 0)
            ws_overview[f'B{row}'] = attendance.get('absent', 0)
            ws_overview[f'C{row}'] = attendance.get('late', 0)
            ws_overview[f'D{row}'] = attendance.get('total', 0)
            
            for col in range(1, 5):
                cell = ws_overview.cell(row=row, column=col)
                cell.font = normal_font
                cell.alignment = center_alignment
                cell.border = thin_border
        
        # Adjust column widths
        ws_overview.column_dimensions['A'].width = 25
        ws_overview.column_dimensions['B'].width = 20
        ws_overview.column_dimensions['C'].width = 15
        ws_overview.column_dimensions['D'].width = 20
        ws_overview.column_dimensions['E'].width = 15
        
        # Sheet 2: Detailed Grades
        if export_options.get('grades', True):
            ws_grades = wb.create_sheet("Chi tiết điểm số")
            
            # Title
            ws_grades['A1'] = 'CHI TIẾT ĐIỂM SỐ'
            ws_grades['A1'].font = title_font
            ws_grades['A1'].alignment = center_alignment
            ws_grades.merge_cells('A1:F1')
            
            # Headers
            grade_headers = ['Bài tập', 'Kỹ năng', 'Điểm', 'Điểm tối đa', 'Ngày nộp', 'Ngày chấm']
            for col, header in enumerate(grade_headers, start=1):
                cell = ws_grades.cell(row=3, column=col)
                cell.value = header
                cell.font = header_font
                cell.fill = PatternFill(start_color='F59E0B', end_color='F59E0B', fill_type='solid')
                cell.alignment = center_alignment
                cell.border = thin_border
            
            # Data
            row = 4
            for grade in progress_data.get('detailed_grades', []):
                ws_grades[f'A{row}'] = grade.get('exercise_title', '')
                ws_grades[f'B{row}'] = grade.get('skill_type') or 'Mixed'
                ws_grades[f'C{row}'] = grade.get('score') if grade.get('score') is not None else 'N/A'
                ws_grades[f'D{row}'] = grade.get('max_score') or 'N/A'
                ws_grades[f'E{row}'] = grade.get('submitted_at', '')[:10] if grade.get('submitted_at') else 'N/A'
                ws_grades[f'F{row}'] = grade.get('graded_at', '')[:10] if grade.get('graded_at') else 'N/A'
                
                for col in range(1, 7):
                    cell = ws_grades.cell(row=row, column=col)
                    cell.font = normal_font
                    cell.border = thin_border
                    if col <= 2 or col >= 5:
                        cell.alignment = left_alignment
                    else:
                        cell.alignment = center_alignment
                
                row += 1
            
            # Adjust column widths
            ws_grades.column_dimensions['A'].width = 40
            ws_grades.column_dimensions['B'].width = 15
            ws_grades.column_dimensions['C'].width = 12
            ws_grades.column_dimensions['D'].width = 15
            ws_grades.column_dimensions['E'].width = 18
            ws_grades.column_dimensions['F'].width = 18
        
        # Sheet 3: Teacher Comments
        if export_options.get('teacherComments', True):
            ws_comments = wb.create_sheet("Nhận xét giáo viên")
            
            # Title
            ws_comments['A1'] = 'NHẬN XÉT CỦA GIÁO VIÊN'
            ws_comments['A1'].font = title_font
            ws_comments['A1'].alignment = center_alignment
            ws_comments.merge_cells('A1:C1')
            
            # Headers
            comment_headers = ['Bài tập', 'Ngày nộp', 'Nhận xét']
            for col, header in enumerate(comment_headers, start=1):
                cell = ws_comments.cell(row=3, column=col)
                cell.value = header
                cell.font = header_font
                cell.fill = PatternFill(start_color='8B5CF6', end_color='8B5CF6', fill_type='solid')
                cell.alignment = center_alignment
                cell.border = thin_border
            
            # Data
            row = 4
            grades_with_feedback = [g for g in progress_data.get('detailed_grades', []) if g.get('feedback')]
            for grade in grades_with_feedback:
                ws_comments[f'A{row}'] = grade.get('exercise_title', '')
                ws_comments[f'B{row}'] = grade.get('submitted_at', '')[:10] if grade.get('submitted_at') else 'N/A'
                ws_comments[f'C{row}'] = grade.get('feedback', 'Không có nhận xét')
                
                for col in range(1, 4):
                    cell = ws_comments.cell(row=row, column=col)
                    cell.font = normal_font
                    cell.alignment = left_alignment if col in [1, 3] else center_alignment
                    cell.border = thin_border
                
                row += 1
            
            if not grades_with_feedback:
                ws_comments['A4'] = 'Chưa có nhận xét từ giáo viên.'
                ws_comments['A4'].font = normal_font
            
            # Adjust column widths
            ws_comments.column_dimensions['A'].width = 40
            ws_comments.column_dimensions['B'].width = 18
            ws_comments.column_dimensions['C'].width = 60
        
        # Save to buffer
        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer


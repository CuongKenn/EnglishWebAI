"""
Lesson Plans Router
API endpoints for lesson plan management and AI generation
"""

from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import get_current_active_user
from app.models.lesson_plan import LessonPlan
from app.models.user import User
from app.schemas.lesson_plan import (
    LessonPlanAIGenerate,
    LessonPlanCreate,
    LessonPlanListResponse,
    LessonPlanResponse,
    LessonPlanUpdate,
)
from app.services.openai_service import openai_service

try:
    from docx import Document
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

router = APIRouter(prefix="/api/v1/lesson-plans", tags=["lesson-plans"])


@router.get("/", response_model=list[LessonPlanListResponse])
async def get_lesson_plans(
    skip: int = 0,
    limit: int = 100,
    grade: int = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Get all lesson plans for the current teacher
    Teachers can only see their own lesson plans
    """
    query = db.query(LessonPlan).filter(LessonPlan.teacher_id == current_user.id)

    if grade:
        query = query.filter(LessonPlan.grade == grade)

    return query.order_by(LessonPlan.created_at.desc()).offset(skip).limit(limit).all()



@router.get("/{lesson_plan_id}", response_model=LessonPlanResponse)
async def get_lesson_plan(
    lesson_plan_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Get a specific lesson plan by ID
    """
    lesson_plan = db.query(LessonPlan).filter(
        LessonPlan.id == lesson_plan_id,
        LessonPlan.teacher_id == current_user.id
    ).first()

    if not lesson_plan:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Lesson plan not found"
        )

    return lesson_plan


@router.post("/", response_model=LessonPlanResponse, status_code=status.HTTP_201_CREATED)
async def create_lesson_plan(
    lesson_plan_data: LessonPlanCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Create a new lesson plan manually
    """
    lesson_plan = LessonPlan(
        teacher_id=current_user.id,
        **lesson_plan_data.dict()
    )

    db.add(lesson_plan)
    db.commit()
    db.refresh(lesson_plan)

    return lesson_plan


@router.post("/generate", response_model=LessonPlanResponse, status_code=status.HTTP_201_CREATED)
async def generate_lesson_plan_with_ai(
    generate_data: LessonPlanAIGenerate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Generate a lesson plan using AI (OpenAI)
    Based on Vietnamese Ngoại ngữ 2018 curriculum
    """
    try:
        # Generate lesson plan with OpenAI
        ai_result = await openai_service.generate_lesson_plan(
            grade=generate_data.grade,
            unit=generate_data.unit,
            lesson_number=generate_data.lesson_number,
            duration=generate_data.duration,
            focus_skills=generate_data.focus_skills,
            language_functions=generate_data.language_functions,
            vocabulary_topics=generate_data.vocabulary_topics,
            grammar_points=generate_data.grammar_points,
            additional_notes=generate_data.additional_notes
        )

        # Create lesson plan from AI result
        lesson_plan = LessonPlan(
            teacher_id=current_user.id,
            title=ai_result.get("title", f"{generate_data.unit} - {generate_data.lesson_number}"),
            subject="English",
            grade=generate_data.grade,
            unit=generate_data.unit,
            lesson_number=generate_data.lesson_number,
            duration=generate_data.duration,
            objectives=ai_result.get("objectives"),
            teaching_aids=ai_result.get("teaching_aids"),
            activities=ai_result.get("activities"),
            content=str(ai_result),  # Store full AI response as content
            notes=ai_result.get("notes"),
            homework=ai_result.get("homework"),
            ai_generated=1,
            ai_prompt=f"Grade: {generate_data.grade}, Unit: {generate_data.unit}, Lesson: {generate_data.lesson_number}"
        )

        db.add(lesson_plan)
        db.commit()
        db.refresh(lesson_plan)

        return lesson_plan

    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate lesson plan: {str(e)}"
        )


@router.put("/{lesson_plan_id}", response_model=LessonPlanResponse)
async def update_lesson_plan(
    lesson_plan_id: int,
    lesson_plan_update: LessonPlanUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Update an existing lesson plan
    """
    lesson_plan = db.query(LessonPlan).filter(
        LessonPlan.id == lesson_plan_id,
        LessonPlan.teacher_id == current_user.id
    ).first()

    if not lesson_plan:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Lesson plan not found"
        )

    # Update fields
    update_data = lesson_plan_update.dict(exclude_unset=True)
    for field, value in update_data.items():
        setattr(lesson_plan, field, value)

    db.commit()
    db.refresh(lesson_plan)

    return lesson_plan


@router.delete("/{lesson_plan_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_lesson_plan(
    lesson_plan_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Delete a lesson plan
    """
    lesson_plan = db.query(LessonPlan).filter(
        LessonPlan.id == lesson_plan_id,
        LessonPlan.teacher_id == current_user.id
    ).first()

    if not lesson_plan:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Lesson plan not found"
        )

    db.delete(lesson_plan)
    db.commit()

    return


@router.get("/{lesson_plan_id}/export/word")
async def export_lesson_plan_word(
    lesson_plan_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Export lesson plan as Word document"""
    if not HAS_DOCX:
        raise HTTPException(
            status_code=status.HTTP_501_NOT_IMPLEMENTED,
            detail="Word export not available. Install python-docx package."
        )

    lesson_plan = db.query(LessonPlan).filter(
        LessonPlan.id == lesson_plan_id,
        LessonPlan.teacher_id == current_user.id
    ).first()

    if not lesson_plan:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Lesson plan not found"
        )

    # Create Word document
    doc = Document()

    # Title
    title = doc.add_heading(lesson_plan.title, level=0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Basic Info
    doc.add_heading('Thông tin cơ bản', level=1)
    doc.add_paragraph(f'Môn học: {lesson_plan.subject}')
    doc.add_paragraph(f'Khối lớp: {lesson_plan.grade}')
    if lesson_plan.unit:
        doc.add_paragraph(f'Unit: {lesson_plan.unit}')
    if lesson_plan.lesson_number:
        doc.add_paragraph(f'Tiết: {lesson_plan.lesson_number}')
    doc.add_paragraph(f'Thời lượng: {lesson_plan.duration} phút')
    doc.add_paragraph('')

    # Objectives
    if lesson_plan.objectives:
        doc.add_heading('Mục tiêu bài học', level=1)
        objectives = lesson_plan.objectives
        if isinstance(objectives, dict):
            for key, items in objectives.items():
                if items and isinstance(items, list):
                    doc.add_heading(key.capitalize(), level=2)
                    for item in items:
                        doc.add_paragraph(f'• {item}')
        doc.add_paragraph('')

    # Teaching Aids
    if lesson_plan.teaching_aids:
        doc.add_heading('Thiết bị và học liệu', level=1)
        for aid in lesson_plan.teaching_aids:
            doc.add_paragraph(f'• {aid}')
        doc.add_paragraph('')

    # Activities
    if lesson_plan.activities:
        doc.add_heading('Tiến trình dạy học', level=1)
        activities = lesson_plan.activities
        if isinstance(activities, dict):
            activity_names = {
                'warm_up': 'Hoạt động 1: Khởi động',
                'presentation': 'Hoạt động 2: Hình thành kiến thức',
                'practice': 'Hoạt động 3: Luyện tập',
                'production': 'Hoạt động 4: Vận dụng'
            }
            for key, name in activity_names.items():
                if key in activities and activities[key]:
                    doc.add_heading(name, level=2)
                    activity = activities[key]
                    if isinstance(activity, dict):
                        if 'content' in activity:
                            doc.add_paragraph(f'Nội dung: {activity["content"]}')
                        if 'duration' in activity:
                            doc.add_paragraph(f'Thời lượng: {activity["duration"]} phút')
                    doc.add_paragraph('')

    # Notes
    if lesson_plan.notes:
        doc.add_heading('Ghi chú', level=1)
        doc.add_paragraph(lesson_plan.notes)
        doc.add_paragraph('')

    # Homework
    if lesson_plan.homework:
        doc.add_heading('Bài tập về nhà', level=1)
        doc.add_paragraph(lesson_plan.homework)

    # Save to BytesIO
    file_stream = BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)

    # Return as downloadable file
    headers = {
        'Content-Disposition': f'attachment; filename="{lesson_plan.title}.docx"'
    }

    return StreamingResponse(
        file_stream,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        headers=headers
    )


@router.get("/{lesson_plan_id}/export/pdf")
async def export_lesson_plan_pdf(
    lesson_plan_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Export lesson plan as PDF (placeholder)"""
    raise HTTPException(
        status_code=status.HTTP_501_NOT_IMPLEMENTED,
        detail="PDF export not yet implemented. Use Word export instead."
    )


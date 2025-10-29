"""
Student Progress Router
API endpoints for student progress tracking, analytics, and reports
"""
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List, Optional
from datetime import datetime
from pydantic import BaseModel
import io

from app.core.database import get_db
from app.core.dependencies import get_current_user
from app.models.user import User, UserRole
from app.models.student_progress import StudentProgressSnapshot, StudentSkillProgress
from app.models.enrollment import Enrollment
from app.models.classroom import Classroom
from app.models.submission import Submission
from app.models.exercise import Exercise
from app.services.student_progress_service import StudentProgressService
from app.services.pdf_report_service import PDFReportService

router = APIRouter(prefix="/api/v1/student-progress", tags=["student-progress"])


# ===================== Schemas =====================

class SkillScore(BaseModel):
    skill: str
    score: float
    count: int


class ProgressSnapshotResponse(BaseModel):
    id: int
    student_id: int
    class_id: Optional[int]
    snapshot_date: datetime
    period_type: str
    period_label: Optional[str]
    total_submissions: int
    graded_submissions: int
    average_score: float
    completion_rate: float
    reading_score: float
    writing_score: float
    listening_score: float
    speaking_score: float
    reading_count: int
    writing_count: int
    listening_count: int
    speaking_count: int
    trend: str
    attendance_rate: float

    class Config:
        from_attributes = True


class SkillProgressResponse(BaseModel):
    id: int
    student_id: int
    skill_type: str
    score: float
    max_score: float
    percentage: float
    assessment_date: datetime
    strengths: Optional[List[str]]
    weaknesses: Optional[List[str]]
    teacher_notes: Optional[str]
    ai_feedback: Optional[str]

    class Config:
        from_attributes = True


class StudentSummaryResponse(BaseModel):
    student_id: int
    latest_snapshot: Optional[dict]
    skill_timelines: dict


class CreateSnapshotRequest(BaseModel):
    student_id: int
    class_id: Optional[int] = None
    period_type: str = "week"
    period_label: Optional[str] = None


# ===================== Endpoints =====================

@router.post("/snapshots", response_model=ProgressSnapshotResponse)
async def create_progress_snapshot(
    request: CreateSnapshotRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Create a progress snapshot for a student
    
    **Permissions:**
    - Teachers can create snapshots for students in their classes
    - Admins can create snapshots for any student
    - Students cannot create snapshots
    """
    # Check permissions
    if current_user.role not in [UserRole.TEACHER, UserRole.ADMIN, UserRole.SUPERADMIN]:
        raise HTTPException(status_code=403, detail="Only teachers and admins can create progress snapshots")
    
    # If teacher, verify they have access to this student
    if current_user.role == UserRole.TEACHER:
        if request.class_id:
            # Check if teacher owns the class
            classroom = db.query(Classroom).filter(Classroom.id == request.class_id).first()
            if not classroom or classroom.teacher_id != current_user.id:
                raise HTTPException(status_code=403, detail="You don't have access to this class")
        else:
            # Check if student is in any of teacher's classes
            enrollment = db.query(Enrollment).join(
                Classroom
            ).filter(
                Enrollment.student_id == request.student_id,
                Classroom.teacher_id == current_user.id
            ).first()
            
            if not enrollment:
                raise HTTPException(status_code=403, detail="Student not in your classes")
    
    # Create snapshot
    snapshot = StudentProgressService.create_progress_snapshot(
        db=db,
        student_id=request.student_id,
        class_id=request.class_id,
        period_type=request.period_type,
        period_label=request.period_label
    )
    
    return snapshot


@router.get("/snapshots/student/{student_id}", response_model=List[ProgressSnapshotResponse])
async def get_student_snapshots(
    student_id: int,
    class_id: Optional[int] = None,
    period_type: str = Query("week", regex="^(week|month|semester)$"),
    limit: int = Query(12, ge=1, le=50),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Get progress snapshots for a student
    
    **Permissions:**
    - Students can view their own snapshots
    - Parents can view their children's snapshots
    - Teachers can view snapshots for students in their classes
    - Admins can view any snapshots
    """
    # Check permissions
    if current_user.role == UserRole.USER:
        # Students can only view their own data
        if student_id != current_user.id:
            raise HTTPException(status_code=403, detail="You can only view your own progress")
    
    elif current_user.role == UserRole.PARENT:
        # Verify this is their child
        from app.models.parent_student import ParentStudent
        link = db.query(ParentStudent).filter(
            ParentStudent.parent_id == current_user.id,
            ParentStudent.student_id == student_id,
            ParentStudent.is_verified == True
        ).first()
        
        if not link:
            raise HTTPException(status_code=403, detail="You can only view your children's progress")
    
    elif current_user.role == UserRole.TEACHER:
        # Verify student is in teacher's class
        enrollment = db.query(Enrollment).join(
            Classroom
        ).filter(
            Enrollment.student_id == student_id,
            Classroom.teacher_id == current_user.id
        ).first()
        
        if not enrollment:
            raise HTTPException(status_code=403, detail="Student not in your classes")
    
    # Get snapshots
    snapshots = StudentProgressService.get_progress_timeline(
        db=db,
        student_id=student_id,
        class_id=class_id,
        period_type=period_type,
        limit=limit
    )
    
    return snapshots


@router.get("/skills/student/{student_id}/{skill_type}", response_model=List[SkillProgressResponse])
async def get_student_skill_timeline(
    student_id: int,
    skill_type: str,
    limit: int = Query(20, ge=1, le=100),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Get skill progress timeline for a student
    
    **Permissions:** Same as get_student_snapshots
    """
    # Check permissions (same logic as above)
    if current_user.role == UserRole.USER:
        if student_id != current_user.id:
            raise HTTPException(status_code=403, detail="You can only view your own progress")
    
    elif current_user.role == UserRole.PARENT:
        from app.models.parent_student import ParentStudent
        link = db.query(ParentStudent).filter(
            ParentStudent.parent_id == current_user.id,
            ParentStudent.student_id == student_id,
            ParentStudent.is_verified == True
        ).first()
        
        if not link:
            raise HTTPException(status_code=403, detail="You can only view your children's progress")
    
    elif current_user.role == UserRole.TEACHER:
        from app.models.classroom import Classroom
        enrollment = db.query(Enrollment).join(
            Classroom
        ).filter(
            Enrollment.student_id == student_id,
            Classroom.teacher_id == current_user.id
        ).first()
        
        if not enrollment:
            raise HTTPException(status_code=403, detail="Student not in your classes")
    
    # Validate skill_type
    if skill_type not in ['reading', 'writing', 'listening', 'speaking']:
        raise HTTPException(status_code=400, detail="Invalid skill type")
    
    # Get skill timeline
    timeline = StudentProgressService.get_skill_timeline(
        db=db,
        student_id=student_id,
        skill_type=skill_type,
        limit=limit
    )
    
    return timeline


@router.get("/summary/student/{student_id}", response_model=StudentSummaryResponse)
async def get_student_summary(
    student_id: int,
    class_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Get comprehensive progress summary for a student
    
    **Permissions:** Same as get_student_snapshots
    """
    # Check permissions
    if current_user.role == UserRole.USER:
        if student_id != current_user.id:
            raise HTTPException(status_code=403, detail="You can only view your own progress")
    
    elif current_user.role == UserRole.PARENT:
        from app.models.parent_student import ParentStudent
        link = db.query(ParentStudent).filter(
            ParentStudent.parent_id == current_user.id,
            ParentStudent.student_id == student_id,
            ParentStudent.is_verified == True
        ).first()
        
        if not link:
            raise HTTPException(status_code=403, detail="You can only view your children's progress")
    
    elif current_user.role == UserRole.TEACHER:
        from app.models.classroom import Classroom
        enrollment = db.query(Enrollment).join(
            Classroom
        ).filter(
            Enrollment.student_id == student_id,
            Classroom.teacher_id == current_user.id
        ).first()
        
        if not enrollment:
            raise HTTPException(status_code=403, detail="Student not in your classes")
    
    # Get summary
    summary = StudentProgressService.get_student_summary(
        db=db,
        student_id=student_id,
        class_id=class_id
    )
    
    return summary


@router.get("/class/{class_id}/snapshots", response_model=List[ProgressSnapshotResponse])
async def get_class_snapshots(
    class_id: int,
    period_type: str = Query("week", regex="^(week|month|semester)$"),
    limit: int = Query(12, ge=1, le=50),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Get all progress snapshots for a class
    
    **Permissions:**
    - Teachers can view snapshots for their classes
    - Admins can view any class snapshots
    """
    if current_user.role not in [UserRole.TEACHER, UserRole.ADMIN, UserRole.SUPERADMIN]:
        raise HTTPException(status_code=403, detail="Only teachers and admins can access this endpoint")
    
    # Verify teacher owns the class
    if current_user.role == UserRole.TEACHER:
        classroom = db.query(Class).filter(Class.id == class_id).first()
        if not classroom or classroom.teacher_id != current_user.id:
            raise HTTPException(status_code=403, detail="You don't have access to this class")
    
    # Get all snapshots for this class
    from sqlalchemy import desc
    snapshots = db.query(StudentProgressSnapshot).filter(
        StudentProgressSnapshot.class_id == class_id,
        StudentProgressSnapshot.period_type == period_type
    ).order_by(
        desc(StudentProgressSnapshot.snapshot_date)
    ).limit(limit * 30).all()  # Get more records for multiple students
    
    return snapshots


@router.get("/export/pdf/student/{student_id}")
async def export_student_report_pdf(
    student_id: int,
    class_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Export comprehensive student progress report as PDF
    
    **Permissions:**
    - Students can export their own reports
    - Parents can export their children's reports
    - Teachers can export reports for students in their classes
    - Admins can export any reports
    
    **Use cases:**
    - Parent-teacher meetings
    - Progress reports for parents
    - Academic records
    """
    # Check permissions (same logic as get_student_snapshots)
    if current_user.role == UserRole.USER:
        if student_id != current_user.id:
            raise HTTPException(status_code=403, detail="You can only export your own report")
    
    elif current_user.role == UserRole.PARENT:
        from app.models.parent_student import ParentStudent
        link = db.query(ParentStudent).filter(
            ParentStudent.parent_id == current_user.id,
            ParentStudent.student_id == student_id,
            ParentStudent.is_verified == True
        ).first()
        
        if not link:
            raise HTTPException(status_code=403, detail="You can only export your children's reports")
    
    elif current_user.role == UserRole.TEACHER:
        enrollment = db.query(Enrollment).join(
            Classroom
        ).filter(
            Enrollment.student_id == student_id,
            Classroom.teacher_id == current_user.id
        ).first()
        
        if not enrollment:
            raise HTTPException(status_code=403, detail="Student not in your classes")
    
    # Get student info
    student = db.query(User).filter(User.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")
    
    # Get class info
    classroom = None
    teacher_name = "N/A"
    class_name = "All Classes"
    
    if class_id:
        classroom = db.query(Classroom).filter(Classroom.id == class_id).first()
        if classroom:
            class_name = classroom.name
            teacher = db.query(User).filter(User.id == classroom.teacher_id).first()
            if teacher:
                teacher_name = teacher.full_name or teacher.username
    else:
        # Get first enrollment
        enrollment = db.query(Enrollment).filter(
            Enrollment.student_id == student_id
        ).first()
        
        if enrollment:
            classroom = db.query(Class).filter(Class.id == enrollment.class_id).first()
            if classroom:
                class_name = classroom.name
                teacher = db.query(User).filter(User.id == classroom.teacher_id).first()
                if teacher:
                    teacher_name = teacher.full_name or teacher.username
    
    # Get latest snapshot
    latest_snapshot = db.query(StudentProgressSnapshot).filter(
        StudentProgressSnapshot.student_id == student_id
    )
    if class_id:
        latest_snapshot = latest_snapshot.filter(StudentProgressSnapshot.class_id == class_id)
    
    latest_snapshot = latest_snapshot.order_by(
        StudentProgressSnapshot.snapshot_date.desc()
    ).first()
    
    overall_score = latest_snapshot.average_score if latest_snapshot else 0.0
    
    # Get skill scores
    skill_scores = {
        'reading': latest_snapshot.reading_score if latest_snapshot else 0.0,
        'writing': latest_snapshot.writing_score if latest_snapshot else 0.0,
        'listening': latest_snapshot.listening_score if latest_snapshot else 0.0,
        'speaking': latest_snapshot.speaking_score if latest_snapshot else 0.0,
    }
    
    # Get progress timeline
    timeline_snapshots = db.query(StudentProgressSnapshot).filter(
        StudentProgressSnapshot.student_id == student_id
    )
    if class_id:
        timeline_snapshots = timeline_snapshots.filter(StudentProgressSnapshot.class_id == class_id)
    
    timeline_snapshots = timeline_snapshots.order_by(
        StudentProgressSnapshot.snapshot_date.asc()
    ).limit(12).all()
    
    progress_timeline = [
        {
            'period_label': s.period_label or s.snapshot_date.strftime('%Y-%m-%d'),
            'average_score': s.average_score,
            'date': s.snapshot_date
        }
        for s in timeline_snapshots
    ]
    
    # Get recent activities
    submissions_query = db.query(Submission).join(
        Exercise
    ).filter(
        Submission.student_id == student_id
    )
    
    if class_id:
        submissions_query = submissions_query.filter(Exercise.class_id == class_id)
    
    recent_submissions = submissions_query.order_by(
        Submission.submitted_at.desc()
    ).limit(10).all()
    
    recent_activities = []
    for sub in recent_submissions:
        if sub.exercise:
            score_pct = None
            if sub.score is not None and sub.exercise.max_score and sub.exercise.max_score > 0:
                score_pct = (sub.score / sub.exercise.max_score) * 100
            
            recent_activities.append({
                'title': sub.exercise.title,
                'date': sub.submitted_at.strftime('%d/%m/%Y') if sub.submitted_at else 'N/A',
                'score': score_pct
            })
    
    # Attendance data (placeholder - would need attendance model)
    attendance_data = {
        'present': 0,
        'absent': 0,
        'late': 0,
        'rate': 0.0
    }
    
    # Teacher comments (could be from latest snapshot or custom)
    teacher_comments = latest_snapshot.notes if latest_snapshot and latest_snapshot.notes else \
                      "Học sinh đang có sự tiến bộ tốt. Tiếp tục duy trì và phát huy."
    
    # Recommendations based on skill scores
    recommendations = []
    for skill, score in skill_scores.items():
        skill_names = {
            'reading': 'kỹ năng đọc',
            'writing': 'kỹ năng viết',
            'listening': 'kỹ năng nghe',
            'speaking': 'kỹ năng nói'
        }
        
        if score < 60:
            recommendations.append(
                f"Cần tăng cường luyện tập {skill_names[skill]}. "
                f"Dành ít nhất 30 phút mỗi ngày để cải thiện."
            )
        elif score < 80:
            recommendations.append(
                f"Tiếp tục phát huy {skill_names[skill]}. "
                f"Có thể thử các bài tập nâng cao hơn."
            )
    
    if not recommendations:
        recommendations.append(
            "Học sinh đang có thành tích xuất sắc. Tiếp tục duy trì và mở rộng vốn từ vựng, "
            "cấu trúc ngữ pháp để nâng cao hơn nữa."
        )
    
    try:
        # Generate PDF
        pdf_buffer = PDFReportService.generate_student_report(
            student_name=student.full_name or student.username,
            student_email=student.email,
            class_name=class_name,
            teacher_name=teacher_name,
            report_date=datetime.now(),
            overall_score=overall_score,
            skill_scores=skill_scores,
            progress_timeline=progress_timeline,
            recent_activities=recent_activities,
            attendance_data=attendance_data,
            teacher_comments=teacher_comments,
            recommendations=recommendations
        )
        
        # Prepare filename
        filename = f"bao_cao_{student.username}_{datetime.now().strftime('%Y%m%d')}.pdf"
        
        return StreamingResponse(
            pdf_buffer,
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Error generating PDF report: {str(e)}"
        )


@router.get("/export/excel/student/{student_id}")
async def export_student_report_excel(
    student_id: int,
    class_id: Optional[int] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Export comprehensive student progress report as Excel
    
    **Permissions:**
    - Students can export their own reports
    - Parents can export their children's reports
    - Teachers can export reports for students in their classes
    - Admins can export any reports
    
    **Use cases:**
    - Parent-teacher meetings
    - Progress reports for parents
    - Academic records
    - Data import into other systems
    """
    # Check permissions (same logic as PDF export)
    if current_user.role == UserRole.USER:
        if student_id != current_user.id:
            raise HTTPException(status_code=403, detail="You can only export your own report")
    
    elif current_user.role == UserRole.PARENT:
        from app.models.parent_student import ParentStudent
        link = db.query(ParentStudent).filter(
            ParentStudent.parent_id == current_user.id,
            ParentStudent.student_id == student_id,
            ParentStudent.is_verified == True
        ).first()
        
        if not link:
            raise HTTPException(status_code=403, detail="You can only export your children's reports")
    
    elif current_user.role == UserRole.TEACHER:
        enrollment = db.query(Enrollment).join(
            Classroom
        ).filter(
            Enrollment.student_id == student_id,
            Classroom.teacher_id == current_user.id
        ).first()
        
        if not enrollment:
            raise HTTPException(status_code=403, detail="Student not in your classes")
    
    # Get student info
    student = db.query(User).filter(User.id == student_id).first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")
    
    # Get class info
    class_name = "All Classes"
    teacher_name = "N/A"
    
    if class_id:
        classroom = db.query(Classroom).filter(Classroom.id == class_id).first()
        if classroom:
            class_name = classroom.name
            teacher = db.query(User).filter(User.id == classroom.teacher_id).first()
            if teacher:
                teacher_name = teacher.full_name or teacher.username
    
    # Get latest snapshot
    latest_snapshot = db.query(StudentProgressSnapshot).filter(
        StudentProgressSnapshot.student_id == student_id
    )
    if class_id:
        latest_snapshot = latest_snapshot.filter(StudentProgressSnapshot.class_id == class_id)
    
    latest_snapshot = latest_snapshot.order_by(
        StudentProgressSnapshot.snapshot_date.desc()
    ).first()
    
    # Get progress timeline
    timeline_snapshots = db.query(StudentProgressSnapshot).filter(
        StudentProgressSnapshot.student_id == student_id
    )
    if class_id:
        timeline_snapshots = timeline_snapshots.filter(StudentProgressSnapshot.class_id == class_id)
    
    timeline_snapshots = timeline_snapshots.order_by(
        StudentProgressSnapshot.snapshot_date.asc()
    ).limit(20).all()
    
    # Get recent submissions
    submissions_query = db.query(Submission).join(
        Exercise
    ).filter(
        Submission.student_id == student_id
    )
    
    if class_id:
        submissions_query = submissions_query.filter(Exercise.class_id == class_id)
    
    recent_submissions = submissions_query.order_by(
        Submission.submitted_at.desc()
    ).limit(20).all()
    
    try:
        # Import openpyxl for Excel generation
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise HTTPException(
                status_code=500,
                detail="Excel export not available. Please contact administrator."
            )
        
        # Create workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Báo cáo tiến bộ"
        
        # Define styles
        header_font = Font(bold=True, size=14, color="FFFFFF")
        header_fill = PatternFill(start_color="667eea", end_color="667eea", fill_type="solid")
        title_font = Font(bold=True, size=16)
        bold_font = Font(bold=True, size=11)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Title
        ws.merge_cells('A1:F1')
        title_cell = ws['A1']
        title_cell.value = "BÁO CÁO TIẾN BỘ HỌC TẬP"
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Student info
        row = 3
        ws[f'A{row}'] = "Học sinh:"
        ws[f'A{row}'].font = bold_font
        ws[f'B{row}'] = student.full_name or student.username
        
        row += 1
        ws[f'A{row}'] = "Email:"
        ws[f'A{row}'].font = bold_font
        ws[f'B{row}'] = student.email
        
        row += 1
        ws[f'A{row}'] = "Lớp:"
        ws[f'A{row}'].font = bold_font
        ws[f'B{row}'] = class_name
        
        row += 1
        ws[f'A{row}'] = "Giáo viên:"
        ws[f'A{row}'].font = bold_font
        ws[f'B{row}'] = teacher_name
        
        row += 1
        ws[f'A{row}'] = "Ngày xuất:"
        ws[f'A{row}'].font = bold_font
        ws[f'B{row}'] = datetime.now().strftime('%d/%m/%Y %H:%M')
        
        # Overall summary
        row += 2
        ws.merge_cells(f'A{row}:F{row}')
        summary_cell = ws[f'A{row}']
        summary_cell.value = "TỔNG QUAN"
        summary_cell.font = header_font
        summary_cell.fill = header_fill
        summary_cell.alignment = Alignment(horizontal="center")
        
        row += 1
        if latest_snapshot:
            headers = ['Điểm TB', 'Xu hướng', 'Bài nộp', 'Bài chấm', 'Tỷ lệ HT']
            for col, header in enumerate(headers, start=1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.font = bold_font
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
            
            row += 1
            trend_text = {
                'improving': 'Đang tiến bộ',
                'declining': 'Cần cải thiện',
                'stable': 'Ổn định'
            }.get(latest_snapshot.trend, 'N/A')
            
            values = [
                f"{latest_snapshot.average_score:.1f}%",
                trend_text,
                latest_snapshot.total_submissions,
                latest_snapshot.graded_submissions,
                f"{latest_snapshot.completion_rate:.1f}%"
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col)
                cell.value = value
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
        
        # Skills breakdown
        row += 2
        ws.merge_cells(f'A{row}:F{row}')
        skills_cell = ws[f'A{row}']
        skills_cell.value = "KẾT QUẢ THEO KỸ NĂNG"
        skills_cell.font = header_font
        skills_cell.fill = header_fill
        skills_cell.alignment = Alignment(horizontal="center")
        
        row += 1
        skill_headers = ['Kỹ năng', 'Điểm', 'Số bài', 'Đánh giá', 'Khuyến nghị']
        for col, header in enumerate(skill_headers, start=1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            cell.font = bold_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center")
        
        if latest_snapshot:
            skills_data = [
                ('Đọc (Reading)', latest_snapshot.reading_score, latest_snapshot.reading_count),
                ('Viết (Writing)', latest_snapshot.writing_score, latest_snapshot.writing_count),
                ('Nghe (Listening)', latest_snapshot.listening_score, latest_snapshot.listening_count),
                ('Nói (Speaking)', latest_snapshot.speaking_score, latest_snapshot.speaking_count)
            ]
            
            for skill_name, score, count in skills_data:
                row += 1
                # Determine evaluation and recommendation
                if score >= 90:
                    evaluation = 'Xuất sắc'
                    recommendation = 'Tiếp tục duy trì'
                elif score >= 80:
                    evaluation = 'Giỏi'
                    recommendation = 'Phát huy thêm'
                elif score >= 70:
                    evaluation = 'Khá'
                    recommendation = 'Cố gắng hơn nữa'
                elif score >= 60:
                    evaluation = 'Trung bình'
                    recommendation = 'Cần luyện tập thêm'
                else:
                    evaluation = 'Cần cải thiện'
                    recommendation = 'Tăng cường luyện tập'
                
                values = [skill_name, f"{score:.1f}%", count, evaluation, recommendation]
                for col, value in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=col)
                    cell.value = value
                    cell.border = border
                    if col == 1:
                        cell.font = bold_font
        
        # Progress timeline
        if timeline_snapshots:
            row += 2
            ws.merge_cells(f'A{row}:F{row}')
            timeline_cell = ws[f'A{row}']
            timeline_cell.value = "TIẾN BỘ THEO THỜI GIAN"
            timeline_cell.font = header_font
            timeline_cell.fill = header_fill
            timeline_cell.alignment = Alignment(horizontal="center")
            
            row += 1
            timeline_headers = ['Thời gian', 'Điểm TB', 'Đọc', 'Viết', 'Nghe', 'Nói']
            for col, header in enumerate(timeline_headers, start=1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.font = bold_font
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
            
            for snapshot in timeline_snapshots:
                row += 1
                values = [
                    snapshot.period_label or snapshot.snapshot_date.strftime('%d/%m/%Y'),
                    f"{snapshot.average_score:.1f}%",
                    f"{snapshot.reading_score:.1f}%",
                    f"{snapshot.writing_score:.1f}%",
                    f"{snapshot.listening_score:.1f}%",
                    f"{snapshot.speaking_score:.1f}%"
                ]
                for col, value in enumerate(values, start=1):
                    cell = ws.cell(row=row, column=col)
                    cell.value = value
                    cell.border = border
                    cell.alignment = Alignment(horizontal="center")
        
        # Recent activities
        if recent_submissions:
            row += 2
            ws.merge_cells(f'A{row}:F{row}')
            activities_cell = ws[f'A{row}']
            activities_cell.value = "BÀI TẬP GẦN ĐÂY"
            activities_cell.font = header_font
            activities_cell.fill = header_fill
            activities_cell.alignment = Alignment(horizontal="center")
            
            row += 1
            activity_headers = ['Tên bài tập', 'Kỹ năng', 'Điểm', 'Ngày nộp', 'Ngày chấm']
            for col, header in enumerate(activity_headers, start=1):
                cell = ws.cell(row=row, column=col)
                cell.value = header
                cell.font = bold_font
                cell.border = border
                cell.alignment = Alignment(horizontal="center")
            
            for sub in recent_submissions:
                if sub.exercise:
                    row += 1
                    score_text = "Chưa chấm"
                    if sub.score is not None and sub.exercise.max_score and sub.exercise.max_score > 0:
                        score_pct = (sub.score / sub.exercise.max_score) * 100
                        score_text = f"{score_pct:.1f}%"
                    
                    values = [
                        sub.exercise.title,
                        sub.exercise.skill_type or 'N/A',
                        score_text,
                        sub.submitted_at.strftime('%d/%m/%Y') if sub.submitted_at else 'N/A',
                        sub.graded_at.strftime('%d/%m/%Y') if sub.graded_at else 'Chưa chấm'
                    ]
                    for col, value in enumerate(values, start=1):
                        cell = ws.cell(row=row, column=col)
                        cell.value = value
                        cell.border = border
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Prepare filename
        filename = f"bao_cao_{student.username}_{datetime.now().strftime('%Y%m%d')}.xlsx"
        
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    
    except ImportError as e:
        raise HTTPException(
            status_code=500,
            detail="Excel export not available. Missing openpyxl library."
        )
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Error generating Excel report: {str(e)}"
        )

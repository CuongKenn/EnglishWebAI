"""
Worksheets Router
API endpoints for worksheet management and AI generation
"""

import logging
import traceback
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.dependencies import get_current_active_user
from app.models.user import User
from app.models.worksheet import Worksheet
from app.schemas.worksheet import (
    WorksheetAIGenerate,
    WorksheetCreate,
    WorksheetListResponse,
    WorksheetResponse,
    WorksheetUpdate,
)
from app.services.openai_service import openai_service

# Setup logging
logger = logging.getLogger(__name__)

try:
    from docx import Document
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    HAS_DOCX = True
    logger.info("python-docx imported successfully")
except ImportError as e:
    HAS_DOCX = False
    logger.error(f"Failed to import python-docx: {e}")

router = APIRouter(prefix="/api/v1/worksheets", tags=["worksheets"])


@router.get("/", response_model=list[WorksheetListResponse])
async def get_worksheets(
    skip: int = 0,
    limit: int = 100,
    grade: int = None,
    worksheet_type: str = None,
    skill_focus: str = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Get all worksheets for the current teacher
    Teachers can only see their own worksheets
    """
    query = db.query(Worksheet).filter(Worksheet.teacher_id == current_user.id)

    if grade:
        query = query.filter(Worksheet.grade == grade)
    if worksheet_type:
        query = query.filter(Worksheet.worksheet_type == worksheet_type)
    if skill_focus:
        query = query.filter(Worksheet.skill_focus == skill_focus)

    return query.order_by(Worksheet.created_at.desc()).offset(skip).limit(limit).all()



@router.get("/{worksheet_id}", response_model=WorksheetResponse)
async def get_worksheet(
    worksheet_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Get a specific worksheet by ID
    """
    worksheet = db.query(Worksheet).filter(
        Worksheet.id == worksheet_id,
        Worksheet.teacher_id == current_user.id
    ).first()

    if not worksheet:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Worksheet not found"
        )

    return worksheet


@router.post("/", response_model=WorksheetResponse, status_code=status.HTTP_201_CREATED)
async def create_worksheet(
    worksheet_data: WorksheetCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Create a new worksheet manually
    """
    worksheet = Worksheet(
        teacher_id=current_user.id,
        **worksheet_data.dict()
    )

    db.add(worksheet)
    db.commit()
    db.refresh(worksheet)

    return worksheet


@router.post("/generate", response_model=WorksheetResponse, status_code=status.HTTP_201_CREATED)
async def generate_worksheet_with_ai(
    generate_data: WorksheetAIGenerate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Generate a worksheet using AI (OpenAI)
    Supports multiple types: multiple_choice, essay, fill_in_blank, etc.
    """
    try:
        # Generate worksheet with OpenAI
        ai_result = await openai_service.generate_worksheet(
            grade=generate_data.grade,
            unit=generate_data.unit,
            worksheet_type=generate_data.worksheet_type,
            skill_focus=generate_data.skill_focus,
            difficulty_level=generate_data.difficulty_level,
            num_questions=generate_data.num_questions,
            duration=generate_data.duration,
            vocabulary_topics=generate_data.vocabulary_topics,
            grammar_points=generate_data.grammar_points,
            language_functions=generate_data.language_functions,
            additional_notes=generate_data.additional_notes
        )

        # Create worksheet from AI result
        worksheet = Worksheet(
            teacher_id=current_user.id,
            title=ai_result.get("title", f"{generate_data.unit} - {generate_data.worksheet_type.replace('_', ' ').title()}"),
            subject="English",
            grade=generate_data.grade,
            unit=generate_data.unit,
            worksheet_type=generate_data.worksheet_type,
            skill_focus=generate_data.skill_focus,
            difficulty_level=generate_data.difficulty_level,
            content=ai_result.get("content"),
            teacher_notes=ai_result.get("teacher_notes"),
            answer_key=ai_result.get("answer_key"),
            duration=generate_data.duration,
            total_points=ai_result.get("total_points", generate_data.num_questions),
            ai_generated=1,
            ai_prompt=f"Grade: {generate_data.grade}, Unit: {generate_data.unit}, Type: {generate_data.worksheet_type}"
        )

        db.add(worksheet)
        db.commit()
        db.refresh(worksheet)

        return worksheet

    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate worksheet: {str(e)}"
        ) from e


@router.put("/{worksheet_id}", response_model=WorksheetResponse)
async def update_worksheet(
    worksheet_id: int,
    worksheet_update: WorksheetUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Update an existing worksheet
    """
    worksheet = db.query(Worksheet).filter(
        Worksheet.id == worksheet_id,
        Worksheet.teacher_id == current_user.id
    ).first()

    if not worksheet:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Worksheet not found"
        )

    # Update fields
    update_data = worksheet_update.dict(exclude_unset=True)
    for field, value in update_data.items():
        setattr(worksheet, field, value)

    db.commit()
    db.refresh(worksheet)

    return worksheet


@router.delete("/{worksheet_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_worksheet(
    worksheet_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """
    Delete a worksheet
    """
    worksheet = db.query(Worksheet).filter(
        Worksheet.id == worksheet_id,
        Worksheet.teacher_id == current_user.id
    ).first()

    if not worksheet:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Worksheet not found"
        )

    db.delete(worksheet)
    db.commit()

    return


@router.get("/{worksheet_id}/export/word")
async def export_worksheet_word(
    worksheet_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Export worksheet as Word document"""
    logger.info(f"Export Word request for worksheet_id={worksheet_id}, user={current_user.id}")

    if not HAS_DOCX:
        logger.error("python-docx not available")
        raise HTTPException(
            status_code=status.HTTP_501_NOT_IMPLEMENTED,
            detail="Word export not available. Install python-docx package."
        )

    worksheet = db.query(Worksheet).filter(
        Worksheet.id == worksheet_id,
        Worksheet.teacher_id == current_user.id
    ).first()

    if not worksheet:
        logger.warning(f"Worksheet {worksheet_id} not found for user {current_user.id}")
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Worksheet not found"
        )

    try:
        logger.info(f"Creating Word document for worksheet: {worksheet.title}")

        # Create Word document
        doc = Document()

        # Title
        title = doc.add_heading(worksheet.title or "Phiếu học tập", level=0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # Basic Info
        doc.add_heading('Thông tin cơ bản', level=1)
        doc.add_paragraph(f'Môn học: {worksheet.subject or "English"}')
        doc.add_paragraph(f'Khối lớp: {worksheet.grade}')
        if worksheet.unit:
            doc.add_paragraph(f'Unit: {worksheet.unit}')
        if worksheet.worksheet_type:
            doc.add_paragraph(f'Loại phiếu: {worksheet.worksheet_type}')
        if worksheet.skill_focus:
            doc.add_paragraph(f'Kỹ năng: {worksheet.skill_focus}')
        if worksheet.difficulty_level:
            doc.add_paragraph(f'Độ khó: {worksheet.difficulty_level}')
        doc.add_paragraph(f'Thời lượng: {worksheet.duration} phút')
        if worksheet.total_points:
            doc.add_paragraph(f'Tổng điểm: {worksheet.total_points}')
        doc.add_paragraph('')

        # Content
        if worksheet.content:
            doc.add_heading('Nội dung', level=1)
            content = worksheet.content

            # Add passage if exists
            if isinstance(content, dict) and 'passage' in content:
                doc.add_heading('Đoạn văn', level=2)
                doc.add_paragraph(content['passage'])
                doc.add_paragraph('')

            if isinstance(content, dict):
                if 'questions' in content:
                    doc.add_heading('Câu hỏi', level=2)
                    for i, q in enumerate(content['questions'], 1):
                        if isinstance(q, dict):
                            # Question number and text
                            question_num = q.get('question_number', i)
                            question_text = q.get('question_text') or q.get('question', '')

                            if question_text:
                                doc.add_paragraph(f'Câu {question_num}: {question_text}', style='Heading 3')
                            else:
                                doc.add_paragraph(f'Câu {question_num}', style='Heading 3')

                            # Options
                            if 'options' in q and isinstance(q['options'], list):
                                for opt in q['options']:
                                    doc.add_paragraph(f'  {opt}')

                            doc.add_paragraph('')
                elif 'text' in content:
                    # For text-based content
                    doc.add_paragraph(str(content['text']))
            elif isinstance(content, str):
                doc.add_paragraph(content)
            doc.add_paragraph('')

        # Teacher Notes
        if worksheet.teacher_notes:
            doc.add_heading('Ghi chú cho giáo viên', level=1)
            doc.add_paragraph(worksheet.teacher_notes)
            doc.add_paragraph('')

        # Answer Key
        if worksheet.answer_key:
            doc.add_heading('Đáp án', level=1)
            answer_key = worksheet.answer_key
            if isinstance(answer_key, dict):
                for key, value in answer_key.items():
                    doc.add_paragraph(f'{key}: {value}')
            elif isinstance(answer_key, str):
                doc.add_paragraph(answer_key)
        elif isinstance(content, dict) and 'questions' in content:
            # Generate answer key from questions if not provided
            doc.add_heading('Đáp án', level=1)
            for q in content['questions']:
                if isinstance(q, dict):
                    question_num = q.get('question_number', '')
                    correct_answer = q.get('correct_answer', '')

                    # If correct_answer is a number, convert to letter (0=A, 1=B, 2=C, 3=D)
                    if isinstance(correct_answer, int):
                        answer_letter = chr(65 + correct_answer)  # 65 is ASCII for 'A'
                        doc.add_paragraph(f'Câu {question_num}: {answer_letter}')
                    else:
                        doc.add_paragraph(f'Câu {question_num}: {correct_answer}')

                    # Add explanation if exists
                    if 'explanation' in q:
                        doc.add_paragraph(f'  Giải thích: {q["explanation"]}')
            doc.add_paragraph('')

        # Save to BytesIO
        logger.info("Saving document to BytesIO")
        file_stream = BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)

        # Return as downloadable file
        # Use URL encoding for non-ASCII characters in filename
        from urllib.parse import quote
        filename = f"{worksheet.title or 'worksheet'}.docx"
        filename = filename.replace('"', '').replace('/', '-').replace('\\', '-')  # Sanitize filename
        filename_encoded = quote(filename)

        headers = {
            'Content-Disposition': f'attachment; filename="worksheet.docx"; filename*=UTF-8\'\'{filename_encoded}'
        }

        logger.info(f"Returning Word document: {filename}")
        return StreamingResponse(
            file_stream,
            media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            headers=headers
        )
    except Exception as e:
        logger.error(f"Error generating Word document: {str(e)}")
        logger.error(traceback.format_exc())
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate Word document: {str(e)}"
        ) from e


@router.get("/{worksheet_id}/export/pdf")
async def export_worksheet_pdf(
    worksheet_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_active_user)
):
    """Export worksheet as PDF (placeholder)"""
    raise HTTPException(
        status_code=status.HTTP_501_NOT_IMPLEMENTED,
        detail="PDF export not yet implemented. Use Word export instead."
    )

